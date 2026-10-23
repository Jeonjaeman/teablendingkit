"""E안(e-editorial.html) → 와디즈 메이커 스튜디오 등록 파일.

  python wadiz/tools/wadiz_export.py [--only story|extra] [--gif-fps 12]
규격: wadiz/완성본파일/KakaoTalk_Longtxt_20261023_1000_24_452.txt (대표 10/23)
출력: wadiz/완성본파일/output/
  01_story/  본문 이미지(JPG, 가로 860, 한 장 세로 3,000px 이하, 10MB 이하)
  02_gif/    움짤(GIF, 가로 860, 5MB 안팎·10MB 이하) + mp4/ (영상 원본)
  03_main/   대표 이미지(16:9, 1200×675 이상, 글자 없음)
  04_intro/  소개 사진(5:3, 800×480 이상, 글자 없음) 최대 10장
  05_reward/ 리워드 이미지(1104×624 이상) 3장
  text.docx · text.md  이미지 사이사이 본문(어느 파일 다음에 어떤 글인지)
  크기검사.txt         파일별 크기·용량·규격 통과 여부

나누는 규칙(전재만 10/23):
  - 글만 있는 단락(제목·본문, 글자에만 모션) → 텍스트 원고
  - 이미지·배경·영상과 모션이 함께 있는 단락 → GIF (글자 모션도 합쳐서)
  - 텍스트로 표현할 수 없는 정적 단락(표·숫자 카드·띠 등) → JPG
  - 표(리워드·원료·레시피)는 와디즈 안내대로 정지 이미지(JPG)
"""
import os, re, io, sys, json, glob, shutil, subprocess, tempfile
from PIL import Image
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from section_gif import serve, SET            # Range 지원 서버 · 애니메이션/영상 시각 맞추기

W = os.path.dirname(HERE)
ROOT = os.path.dirname(W)
OUT = os.path.join(W, '완성본파일', 'output')
MB = 1024 * 1024
GIF_SOFT, GIF_HARD, IMG_HARD, MAX_H = 5 * MB, 10 * MB, 10 * MB, 3000
PAD = 14

# 내보내기용 레이아웃: 왼쪽 라벨 열(lead)은 텍스트 원고로 빼고, 본문을 전체 폭으로
EXPORT_CSS = '''
.sec>.lead{display:none!important}
.sec>.body{grid-column:1/-1!important}
.mast{display:none!important}
'''

# 단락 나누기: 텍스트 단위 / 화면 단위를 페이지 순서대로
SCAN = r'''() => {
  const TXT = el => {
    const t = el.tagName.toLowerCase(), c = el.classList;
    if (['h1','h2','h3','p','dl'].includes(t)) return true;
    if (t === 'ul' && (c.contains('says') || c.contains('fine'))) return true;
    if (t === 'ol' && (c.contains('rows') || c.contains('parts'))) return true;
    if (t === 'div' && (c.contains('ts') || c.contains('links'))) return true;
    return false;
  };
  const lines = el => (el.innerText || '').split('\n').map(s => s.trim()).filter(Boolean);
  const textOf = el => {
    const t = el.tagName.toLowerCase(), c = el.classList, out = [];
    if (t === 'dl') { el.querySelectorAll('dt').forEach(dt => { out.push(['q', dt.innerText.trim()]);
        const dd = dt.nextElementSibling; if (dd) out.push(['a', dd.innerText.trim()]); }); return out; }
    if (t === 'ul' || t === 'ol') { [...el.querySelectorAll(':scope > li')].sort((x, y) => c.contains('parts')   // 2열 목록은 번호순
        ? (parseInt(x.innerText) || 0) - (parseInt(y.innerText) || 0) : 0).forEach(li => {
        const s = li.innerText.replace(/\s*\n\s*/g, ' ').trim(); if (s) out.push([c.contains('says') ? 'quote' : 'bullet', s]); }); return out; }
    const style = /^h[123]$/.test(t) ? t : c.contains('sub') ? 'lead' : c.contains('cap') ? 'cap' : c.contains('quote') ? 'quote' : 'p';
    const s = (el.innerText || '').trim(); if (s) out.push([style, s]);
    return out;
  };
  const motionNonText = el => [...el.querySelectorAll('*')].concat([el]).some(e => {
    if (e.closest('[data-mo]')) return false;                      // 글자 모션은 따로
    return [...e.classList].some(k => k.startsWith('mo-') && !/^mo-(tx|w|u|l|d\d)/.test(k));
  });
  const units = []; let n = 0;
  const vis = (el, sec) => {
    el.dataset.xid = 'x' + (n++);
    const vids = [...el.querySelectorAll('video')];
    const hasImg = el.querySelector('img, svg, video') !== null;
    const isTable = el.tagName === 'TABLE' || el.querySelector('table') !== null;
    const textMo = el.querySelector('[data-mo]') !== null;
    const gif = !isTable && (vids.length > 0 || motionNonText(el) || (textMo && hasImg));
    const name = vids.length ? (vids[0].querySelector('source')?.getAttribute('src') || '').split('/').pop().replace('.mp4', '')
      : (el.className && typeof el.className === 'string' ? el.className.split(' ').filter(k => !['full','clip','media','grid2','grid3'].includes(k))[0] : '')
        || el.tagName.toLowerCase();
    units.push({kind: 'vis', xid: el.dataset.xid, sec, gif, table: isTable, name,
      videos: vids.map(v => (v.querySelector('source')?.getAttribute('src') || ''))});
  };
  for (const s of document.querySelectorAll('section.sec')) {
    const sec = s.id;
    for (const c of s.children) {
      if (c.classList.contains('lead')) {
        // 내보낼 때 lead 를 숨기므로 innerText 에 줄바꿈이 없다 → <br> 을 공백으로 직접 바꿈
        const tx = e => e ? e.innerHTML.replace(/<br\s*\/?>/gi, ' ').replace(/<[^>]+>/g, '').replace(/\s+/g, ' ').trim() : '';
        const num = tx(c.querySelector('.num')), tag = tx(c.querySelector('.tag')), en = tx(c.querySelector('.en'));
        if (num || tag || en) units.push({kind: 'text', sec, items: [['label', [num, tag].filter(Boolean).join(' · ') + (en ? ' — ' + en : '')]]});
        continue;
      }
      if (c.classList.contains('body')) {
        for (const k of c.children) {
          if (TXT(k) && !k.querySelector('img, svg, video')) units.push({kind: 'text', sec, items: textOf(k)});
          else vis(k, sec);
        }
        continue;
      }
      vis(c, sec);
    }
  }
  const f = document.querySelector('footer');
  if (f) units.push({kind: 'text', sec: 'footer', items: [['cap', f.innerText.replace(/\s*\n\s*/g, ' ').trim()]]});
  return units;
}'''


def bbox(pg, xid):
    return pg.evaluate('''x => { const e = document.querySelector(`[data-xid="${x}"]`); const r = e.getBoundingClientRect();
        return {top: r.top + scrollY, bottom: r.bottom + scrollY}; }''', xid)


def size_ok_jpg(path):
    im = Image.open(path)
    return im.width, im.height, os.path.getsize(path)


def save_jpg(im, path):
    for q in (90, 85, 80, 72, 64):
        im.convert('RGB').save(path, quality=q, optimize=True, progressive=True)
        if os.path.getsize(path) <= IMG_HARD:
            return q
    raise RuntimeError('JPG too big ' + path)


def capture_static(pg, top, bottom, path_base):
    """페이지 [top, bottom] 구간을 가로 860 전체로 찍어 3,000px 넘으면 나눔."""
    top, bottom = max(0, int(top - PAD)), int(bottom + PAD)
    shot = pg.screenshot(full_page=True, clip={'x': 0, 'y': top, 'width': 860, 'height': bottom - top})
    im = Image.open(io.BytesIO(shot))
    paths = []
    parts = max(1, -(-im.height // MAX_H))
    step = -(-im.height // parts)
    for i in range(parts):
        seg = im.crop((0, i * step, 860, min(im.height, (i + 1) * step)))
        p = path_base + ('' if parts == 1 else f'_{i + 1}') + '.jpg'
        save_jpg(seg, p)
        paths.append(p)
    return paths


LADDER = [  # (프레임 간격, 색 수, 디더)
    (1, 256, 'sierra2_4a'), (1, 192, 'bayer:bayer_scale=3'), (1, 128, 'bayer:bayer_scale=3'),
    (2, 256, 'sierra2_4a'), (2, 160, 'bayer:bayer_scale=3'), (2, 96, 'bayer:bayer_scale=4'),
    (3, 128, 'bayer:bayer_scale=4'), (3, 64, 'bayer:bayer_scale=5'),
]


def encode_gif(frames_dir, n, fps, out):
    """용량 사다리: 5MB 이하가 나오는 첫 조합, 없으면 10MB 이하 중 가장 작은 것."""
    tried = []
    for step, colors, dither in LADDER:
        tmp = out + '.try.gif'
        pre = f"select='not(mod(n\\,{step}))',setpts=N/({fps / step})/TB" if step > 1 else 'null'
        pal = os.path.join(frames_dir, 'pal.png')
        src = ['-framerate', str(fps), '-i', os.path.join(frames_dir, '%04d.png')]
        subprocess.run(['ffmpeg', '-v', 'error', '-y', *src, '-vf', f'{pre},palettegen=max_colors={colors}:stats_mode=diff', pal], check=True)
        r = subprocess.run(['ffmpeg', '-v', 'error', '-y', *src, '-i', pal,
                            '-lavfi', f'[0:v]{pre}[s];[s][1:v]paletteuse=dither={dither}:diff_mode=rectangle',
                            '-r', f'{fps / step:g}', '-loop', '0', tmp], capture_output=True)
        if r.returncode != 0:
            continue
        sz = os.path.getsize(tmp)
        tried.append((sz, step, colors, dither))
        if sz <= GIF_SOFT:
            shutil.move(tmp, out)
            for f in glob.glob(out + '.*.gif'):
                os.remove(f)
            return sz, fps / step, colors
        shutil.copy(tmp, out + f'.{len(tried)}.gif')
    ok = [t for t in tried if t[0] <= GIF_HARD]
    if not ok:
        raise RuntimeError('GIF too big even at lowest setting: ' + out)
    best = min(ok)
    i = tried.index(best) + 1
    shutil.move(out + f'.{i}.gif', out)
    for f in glob.glob(out + '.*.gif') + glob.glob(out + '.try.gif'):
        os.remove(f)
    return best[0], fps / best[1], best[2]


def capture_gif(pg, xid, dur, fps, out):
    pg.evaluate("x => document.querySelector(`[data-xid=\"${x}\"]`).scrollIntoView({block: 'center'})", xid)
    pg.wait_for_timeout(300)
    pg.evaluate(SET, 0)
    r = pg.evaluate('''x => { const e = document.querySelector(`[data-xid="${x}"]`).getBoundingClientRect(); return {t: e.top, b: e.bottom}; }''', xid)
    top = max(0, int(r['t'] - PAD)); h = int(r['b'] + PAD) - top
    clip = {'x': 0, 'y': top, 'width': 860, 'height': h}
    tmp = tempfile.mkdtemp()
    try:
        n = int(round(dur * fps))
        for i in range(n):
            pg.evaluate(SET, i / fps)
            Image.open(io.BytesIO(pg.screenshot(clip=clip))).convert('RGB').save(os.path.join(tmp, f'{i:04d}.png'))
        for f in glob.glob(out + '*'):
            os.remove(f)
        return encode_gif(tmp, n, fps, out)
    finally:
        shutil.rmtree(tmp)


def story(fps, text_only=False):
    for d in ('01_story', '02_gif', '02_gif/mp4'):
        os.makedirs(os.path.join(OUT, d), exist_ok=True)
    for f in [] if text_only else glob.glob(os.path.join(OUT, '01_story', '*')) + glob.glob(os.path.join(OUT, '02_gif', '*.gif')) + glob.glob(os.path.join(OUT, '02_gif', 'mp4', '*')):
        os.remove(f)
    srv = serve(); port = srv.server_address[1]
    report, doc = [], []
    try:
        with sync_playwright() as p:
            br = p.chromium.launch()
            pg = br.new_page(viewport={'width': 860, 'height': 1400}, device_scale_factor=1)
            pg.goto(f'http://127.0.0.1:{port}/wadiz/e-editorial.html', wait_until='load')
            pg.add_style_tag(content=EXPORT_CSS)
            pg.evaluate("document.getAnimations().forEach(a => a.pause())")
            pg.evaluate('''() => Promise.all([...document.querySelectorAll('video')].map(v => { v.preload = 'auto'; v.removeAttribute('autoplay'); v.pause();
                return v.readyState >= 2 ? 0 : new Promise(r => { v.addEventListener('loadeddata', r, {once: true}); v.load(); setTimeout(r, 8000); }); }))''')
            pg.wait_for_timeout(500)
            units = pg.evaluate(SCAN)
            durs = pg.evaluate('''() => Object.fromEntries([...document.querySelectorAll('[data-xid]')].map(e =>
                [e.dataset.xid, Math.max(0, ...[...e.querySelectorAll('video')].map(v => v.duration || 0))]))''')
            # 정적 단위끼리 연속이면(같은 단락) 한 장으로 합침
            merged = []
            for u in units:
                if (u['kind'] == 'vis' and not u['gif'] and merged and merged[-1]['kind'] == 'vis' and not merged[-1]['gif']
                        and merged[-1]['sec'] == u['sec']):
                    merged[-1]['xids'].append(u['xid']); merged[-1]['name'] += '+' + u['name']
                elif u['kind'] == 'vis':
                    merged.append(dict(u, xids=[u['xid']]))
                elif merged and merged[-1]['kind'] == 'text':
                    merged[-1]['items'] += u['items']
                else:
                    merged.append(dict(u))
            pg.evaluate(SET, 0)
            idx = 0
            if text_only:          # 캡처 없이 원고만 다시 — 이미 만든 파일 번호에 맞춤
                for u in merged:
                    if u['kind'] == 'text':
                        doc.append(('text', u['items'])); continue
                    idx += 1
                    for f in sorted(glob.glob(os.path.join(OUT, '01_story', f'{idx:02d}_*')) + glob.glob(os.path.join(OUT, '02_gif', f'{idx:02d}_*.gif'))):
                        doc.append(('file', os.path.relpath(f, OUT)))
                merged = []
            for u in merged:
                if u['kind'] == 'text':
                    doc.append(('text', u['items']))
                    continue
                idx += 1
                slug = re.sub(r'[^0-9A-Za-z_]+', '_', u['name'])[:28].strip('_')
                base = f'{idx:02d}_{u["sec"]}_{slug}'
                if u['gif']:
                    dur = max(4.0, durs.get(u['xid'], 0) or 0)
                    out = os.path.join(OUT, '02_gif', base + '.gif')
                    sz, f2, colors = capture_gif(pg, u['xid'], dur, fps, out)
                    im = Image.open(out)
                    report.append((os.path.relpath(out, OUT), im.width, im.height, sz, f'{dur:g}s {f2:g}fps {colors}색'))
                    doc.append(('file', os.path.relpath(out, OUT)))
                    for v in u['videos']:
                        src = os.path.join(W, v.replace('/', os.sep))
                        if os.path.exists(src):
                            dst = os.path.join(OUT, '02_gif', 'mp4', base + '.mp4')
                            shutil.copy(src, dst)
                    pg.evaluate(SET, 0)
                else:
                    boxes = [bbox(pg, x) for x in u['xids']]
                    paths = capture_static(pg, min(b['top'] for b in boxes), max(b['bottom'] for b in boxes), os.path.join(OUT, '01_story', base))
                    for pth in paths:
                        w, h, sz = size_ok_jpg(pth)
                        report.append((os.path.relpath(pth, OUT), w, h, sz, '표' if u['table'] else ''))
                        doc.append(('file', os.path.relpath(pth, OUT)))
            br.close()
    finally:
        srv.shutdown()
    write_text(doc)
    return report


STYLE_KO = {'label': '단락 표시', 'h1': '제목', 'h2': '제목', 'h3': '소제목', 'lead': '본문(강조)', 'p': '본문',
            'cap': '작은 글씨', 'quote': '인용', 'bullet': '목록', 'q': '질문', 'a': '답변'}


def write_text(doc):
    import docx
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Malgun Gothic'; st.font.size = Pt(10.5)
    d.add_heading('취향의 권리 — 와디즈 스토리 본문 원고', level=1)
    d.add_paragraph('순서대로 올려 주세요. [파일] 은 이미지·GIF, 그 사이 글은 에디터 텍스트로 넣습니다. '
                    '(단락 표시)는 섹션 머리 문구로, 작게 넣거나 생략해도 됩니다.')
    md = ['# 취향의 권리 — 와디즈 스토리 본문 원고', '']
    for kind, payload in doc:
        if kind == 'file':
            p = d.add_paragraph(); r = p.add_run(f'[{payload.replace(os.sep, "/")}]'); r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4D, 0x3A)
            md += [f'**[{payload.replace(os.sep, "/")}]**', '']
            continue
        for style, text in payload:
            if style in ('h1', 'h2'):
                d.add_heading(text, level=2); md += ['## ' + text.replace('\n', ' '), '']
            elif style == 'h3':
                d.add_heading(text, level=3); md += ['### ' + text.replace('\n', ' '), '']
            elif style == 'label':
                p = d.add_paragraph(); r = p.add_run(text); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                md += [f'<sub>{text}</sub>', '']
            elif style == 'bullet':
                d.add_paragraph(text, style='List Bullet'); md += ['- ' + text]
            elif style == 'q':
                p = d.add_paragraph(); p.add_run(text).bold = True; md += ['**' + text + '**', '']
            elif style == 'quote':
                p = d.add_paragraph(); r = p.add_run(text); r.italic = True; md += ['> ' + text.replace('\n', ' '), '']
            else:
                p = d.add_paragraph()
                r = p.add_run(text)
                if style == 'lead': r.bold = True
                if style == 'cap': r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                md += [text.replace('\n', '  \n'), '']
        md.append('')
    d.save(os.path.join(OUT, 'text.docx'))
    open(os.path.join(OUT, 'text.md'), 'w', encoding='utf-8').write('\n'.join(md))


# ---------- 대표·소개·리워드 ----------
def crop_to(im, ratio):
    w, h = im.size
    if w / h > ratio:
        nw = int(h * ratio); x = (w - nw) // 2; return im.crop((x, 0, x + nw, h))
    nh = int(w / ratio); y = (h - nh) // 2; return im.crop((0, y, w, y + nh))


def video_frame(name, t):
    src = os.path.join(W, 'assets', 'loop', name + '.mp4')
    tmp = tempfile.mktemp(suffix='.png')
    subprocess.run(['ffmpeg', '-v', 'error', '-y', '-ss', str(t), '-i', src, '-frames:v', '1', tmp], check=True)
    im = Image.open(tmp).convert('RGB'); os.remove(tmp)
    return im


def extra():
    rep = []
    for d in ('03_main', '04_intro', '05_reward'):
        os.makedirs(os.path.join(OUT, d), exist_ok=True)
        for f in glob.glob(os.path.join(OUT, d, '*')):
            os.remove(f)
    # 대표 이미지(16:9, 글자 없음): 실사 원료 6종(VER.1) — 대안: 찻잔에 물 붓는 장면
    main = crop_to(Image.open(os.path.join(W, 'assets', 'real', 'herbs_taste.jpg')), 16 / 9).resize((1500, 844), Image.LANCZOS)
    p = os.path.join(OUT, '03_main', 'main_1_추천_원료6종.jpg'); save_jpg(main, p); rep.append(p)
    alt = crop_to(video_frame('c09_pour_water', 1.0), 16 / 9).resize((1280, 720), Image.LANCZOS)
    p = os.path.join(OUT, '03_main', 'main_2_대안_물붓기.jpg'); save_jpg(alt, p); rep.append(p)
    # 소개 사진(5:3, 글자 없음) — 실사 원료 + 라이프스타일 영상 정지 컷
    intro = [('herbs_taste', Image.open(os.path.join(W, 'assets', 'real', 'herbs_taste.jpg'))),
             ('herbs_depth', Image.open(os.path.join(W, 'assets', 'real', 'herbs_depth.jpg')))]
    for name, t in [('c09_pour_water', 1.0), ('c04_pour_cup', 1.0), ('c06_smell', 1.0), ('c11_brew_closeup', 1.0),
                    ('c01_teapot', 2.0), ('c12_window_sit', 2.0), ('c13_smile', 2.0), ('c10_rest', 2.0)]:
        intro.append((name, video_frame(name, t)))
    for i, (name, im) in enumerate(intro[:10], 1):
        im = crop_to(im.convert('RGB'), 5 / 3)
        im = im.resize((1200, 720), Image.LANCZOS)
        p = os.path.join(OUT, '04_intro', f'intro_{i:02d}_{name}.jpg'); save_jpg(im, p); rep.append(p)
    # 리워드(1104×624 이상): 블렌더 연출 컷 — 글자 없는 단순 배경(제품 인쇄면은 실물 그대로)
    for i, (name, f) in enumerate([('울트라얼리버드', 'flatlay_ultra'), ('슈퍼얼리버드', 'flatlay_super'), ('2개세트', 'flatlay_set')], 1):
        src = os.path.join(ROOT, '3d', 'render', f + '.png')
        if not os.path.exists(src):
            src = os.path.join(ROOT, '3d', 'render', f + '.jpg')
        im = crop_to(Image.open(src).convert('RGB'), 1104 / 624).resize((1656, 936), Image.LANCZOS)
        p = os.path.join(OUT, '05_reward', f'reward_{i}_{name}.jpg'); save_jpg(im, p); rep.append(p)
    out = []
    for p in rep:
        im = Image.open(p); out.append((os.path.relpath(p, OUT), im.width, im.height, os.path.getsize(p), ''))
    return out


def check(report):
    """규격 검사 → 크기검사.txt"""
    lines = ['파일 | 가로×세로 | 용량 | 비고 | 판정']
    bad = []
    for rel, w, h, sz, note in report:
        folder = rel.split(os.sep)[0]
        if folder == '01_story':
            ok = w == 860 and h <= MAX_H and sz <= IMG_HARD
        elif folder == '02_gif':
            ok = w == 860 and sz <= GIF_HARD
            note += ' (5MB 이하 권장 충족)' if sz <= GIF_SOFT else ' (5MB 초과·10MB 이하)'
        elif folder == '03_main':
            ok = w >= 1200 and h >= 675 and abs(w / h - 16 / 9) < .01 and sz <= IMG_HARD
        elif folder == '04_intro':
            ok = w >= 800 and h >= 480 and sz <= IMG_HARD
        elif folder == '05_reward':
            ok = w >= 1104 and h >= 624 and sz <= IMG_HARD
        else:
            ok = True
        lines.append(f'{rel} | {w}×{h} | {sz / MB:.2f}MB | {note} | {"통과" if ok else "불합격"}')
        if not ok:
            bad.append(rel)
    return lines, bad


def main():
    only = sys.argv[sys.argv.index('--only') + 1] if '--only' in sys.argv else None
    fps = int(sys.argv[sys.argv.index('--gif-fps') + 1]) if '--gif-fps' in sys.argv else 12
    os.makedirs(OUT, exist_ok=True)
    rep = []
    if only in (None, 'story'):
        rep += story(fps)
    if only == 'text':
        story(fps, text_only=True)
        print('text ok')
        return
    if only in (None, 'extra'):
        rep += extra()
    lines, bad = check(rep)
    open(os.path.join(OUT, '크기검사.txt' if only is None else f'크기검사_{only}.txt'), 'w', encoding='utf-8').write('\n'.join(lines))
    print('\n'.join(lines))
    print('불합격', len(bad), bad)


if __name__ == '__main__':
    main()
