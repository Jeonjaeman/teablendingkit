"""E안 → 와디즈 등록 파일 2안: 본문을 전부 이미지·GIF로(글자 포함), 맨 아래에 검색용 텍스트.

  python wadiz/tools/wadiz_export_full.py
출력: wadiz/완성본파일/output2/
  01_story/  정지 줄(JPG, 가로 860, 한 장 세로 3,000px 이하)
  02_gif/    모션이 있는 줄(GIF, 가로 860, 5MB 안팎·10MB 이하) + mp4/
  03_main · 04_intro · 05_reward  (output/ 과 같은 파일을 복사)
  text.docx · text.md  ① 올리는 순서(파일 번호) ② 맨 아래에 넣을 검색용 텍스트(페이지 전체 글)
  크기검사.txt

자르는 규칙(전재만 10/23):
  - 페이지를 '줄' 단위로 자른다. 같은 높이에 나란히 있는 칸(왼쪽 번호 열 + 본문, 2·3칸 그리드)은 한 장.
  - 줄 사이 여백의 가운데에서 잘라, 순서대로 쌓으면 원래 페이지와 똑같이 이어진다.
  - 줄 안에 영상이나 모션(글자 모션 포함)이 있으면 GIF, 없으면 JPG. 이웃한 정지 줄은 3,000px까지 한 장으로 합친다.
"""
import os, io, sys, glob, shutil, subprocess, tempfile
from PIL import Image
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from section_gif import serve, SET
import wadiz_export as X            # encode_gif(용량 사다리) · save_jpg · check · 규격 상수

W = X.W
OUT = os.path.join(W, '완성본파일', 'output2')
OUT1 = X.OUT

ROWS = r'''() => {
  const motion = el => el.querySelector('video, [data-mo], [class*="mo-"]') !== null || [...el.classList].some(k => k.startsWith('mo-'));
  const dur = el => Math.max(0, ...[...el.querySelectorAll('video')].map(v => v.duration || 0));
  const blocks = [];
  for (const s of document.querySelectorAll('header.mast, section.sec, footer')) {
    const kids = s.tagName === 'SECTION' ? [...s.children] : [s];
    const rs = kids.map(k => ({k, r: k.getBoundingClientRect()})).filter(o => o.r.height > 0);
    // 세로 구간이 겹치는 칸끼리 한 줄
    const groups = [];
    for (const o of rs.sort((a, b) => a.r.top - b.r.top)) {
      const g = groups[groups.length - 1];
      if (g && o.r.top < g.bottom - 1) { g.bottom = Math.max(g.bottom, o.r.bottom); g.els.push(o.k); }
      else groups.push({top: o.r.top, bottom: o.r.bottom, els: [o.k]});
    }
    for (const g of groups) blocks.push({sec: s.id || s.tagName.toLowerCase(), top: g.top + scrollY, bottom: g.bottom + scrollY,
      motion: g.els.some(motion), dur: Math.max(0, ...g.els.map(dur)),
      vids: g.els.flatMap(e => [...e.querySelectorAll('video source')].map(x => x.getAttribute('src')))});
  }
  return {blocks, H: document.documentElement.scrollHeight};
}'''

TEXT = r'''() => {
  const out = [];
  for (const s of document.querySelectorAll('section.sec')) {
    const head = [s.querySelector('.lead .num'), s.querySelector('.lead .tag')].filter(Boolean).map(e => e.innerText.trim()).join(' · ');
    const lines = (s.innerText || '').split('\n').map(x => x.trim()).filter(Boolean);
    out.push({head, lines});
  }
  return out;
}'''


def frames(pg, top, h, dur, fps, tmp):
    pg.evaluate("y => window.scrollTo(0, y)", max(0, top - 200))
    pg.wait_for_timeout(200)
    off = pg.evaluate('scrollY')
    clip = {'x': 0, 'y': top - off, 'width': 860, 'height': h}
    n = int(round(dur * fps))
    for i in range(n):
        pg.evaluate(SET, i / fps)
        Image.open(io.BytesIO(pg.screenshot(clip=clip))).convert('RGB').save(os.path.join(tmp, f'{i:04d}.png'))
    return n


def main(fps=12):
    for d in ('01_story', '02_gif', '02_gif/mp4'):
        os.makedirs(os.path.join(OUT, d), exist_ok=True)
        for f in glob.glob(os.path.join(OUT, d, '*.*')):
            os.remove(f)
    srv = serve(); port = srv.server_address[1]
    report, order = [], []
    try:
        with sync_playwright() as p:
            br = p.chromium.launch()
            pg = br.new_page(viewport={'width': 860, 'height': 3400}, device_scale_factor=1)
            pg.goto(f'http://127.0.0.1:{port}/wadiz/e-editorial.html', wait_until='load')
            pg.evaluate("document.getAnimations().forEach(a => a.pause())")
            pg.evaluate('''() => Promise.all([...document.querySelectorAll('video')].map(v => { v.preload = 'auto'; v.removeAttribute('autoplay'); v.pause();
                return v.readyState >= 2 ? 0 : new Promise(r => { v.addEventListener('loadeddata', r, {once: true}); v.load(); setTimeout(r, 8000); }); }))''')
            pg.wait_for_timeout(500)
            pg.evaluate(SET, 0)
            info = pg.evaluate(ROWS)
            texts = pg.evaluate(TEXT)
            B, H = info['blocks'], info['H']
            # 줄 사이 여백 가운데에서 자름 → 빈틈 없이 페이지 전체
            cuts = [0] + [round((B[i]['bottom'] + B[i + 1]['top']) / 2) for i in range(len(B) - 1)] + [H]
            pieces = [dict(B[i], y0=cuts[i], y1=cuts[i + 1]) for i in range(len(B))]
            # 이웃 정지 줄 합치기(같은 섹션, 3,000px 이하)
            merged = []
            for pc in pieces:
                m = merged[-1] if merged else None
                if m and not m['motion'] and not pc['motion'] and m['sec'] == pc['sec'] and pc['y1'] - m['y0'] <= X.MAX_H:
                    m['y1'] = pc['y1']
                else:
                    merged.append(dict(pc))
            full = Image.open(io.BytesIO(pg.screenshot(full_page=True)))
            for i, pc in enumerate(merged, 1):
                base = f"{i:02d}_{pc['sec']}"
                h = pc['y1'] - pc['y0']
                if not pc['motion']:
                    im = full.crop((0, pc['y0'], 860, pc['y1']))
                    parts = max(1, -(-h // X.MAX_H)); step = -(-h // parts)
                    for k in range(parts):
                        pth = os.path.join(OUT, '01_story', base + ('' if parts == 1 else f'_{k + 1}') + '.jpg')
                        X.save_jpg(im.crop((0, k * step, 860, min(h, (k + 1) * step))), pth)
                        order.append(os.path.relpath(pth, OUT))
                        w_, h_ = Image.open(pth).size
                        report.append((os.path.relpath(pth, OUT), w_, h_, os.path.getsize(pth), ''))
                    continue
                dur = max(4.0, pc['dur'])
                out = os.path.join(OUT, '02_gif', base + '.gif')
                tmp = tempfile.mkdtemp()
                try:
                    n = frames(pg, pc['y0'], h, dur, fps, tmp)
                    sz, f2, colors = X.encode_gif(tmp, n, fps, out)
                finally:
                    shutil.rmtree(tmp)
                pg.evaluate(SET, 0)
                w_, h_ = Image.open(out).size
                report.append((os.path.relpath(out, OUT), w_, h_, sz, f'{dur:g}s {f2:g}fps {colors}색'))
                order.append(os.path.relpath(out, OUT))
                for v in pc['vids']:
                    src = os.path.join(W, v.replace('/', os.sep))
                    if os.path.exists(src):
                        shutil.copy(src, os.path.join(OUT, '02_gif', 'mp4', f"{base}_{os.path.basename(v)}"))
            br.close()
    finally:
        srv.shutdown()
    # 대표·소개·리워드: 1안과 같은 파일
    for d in ('03_main', '04_intro', '05_reward'):
        if os.path.isdir(os.path.join(OUT, d)):
            shutil.rmtree(os.path.join(OUT, d))
        shutil.copytree(os.path.join(OUT1, d), os.path.join(OUT, d))
        for f in sorted(glob.glob(os.path.join(OUT, d, '*'))):
            w_, h_ = Image.open(f).size
            report.append((os.path.relpath(f, OUT), w_, h_, os.path.getsize(f), ''))
    write_text(order, texts)
    lines, bad = X.check(report)
    open(os.path.join(OUT, '크기검사.txt'), 'w', encoding='utf-8').write(
        '와디즈 등록 파일 규격 검사 — 2안(본문 전부 이미지·GIF)\n\n' + '\n'.join(lines))
    print('\n'.join(lines)); print('불합격', len(bad), bad)


def write_text(order, texts):
    import docx
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    d.styles['Normal'].font.name = 'Malgun Gothic'; d.styles['Normal'].font.size = Pt(10.5)
    d.add_heading('취향의 권리 — 와디즈 스토리 등록 원고 (2안)', level=1)
    d.add_paragraph('본문은 아래 순서대로 이미지·GIF만 이어 붙입니다. 글자는 모두 이미지 안에 들어 있습니다. '
                    '마지막에 "검색용 텍스트"를 에디터 텍스트로 넣어 주세요(검색·접근성용).')
    md = ['# 취향의 권리 — 와디즈 스토리 등록 원고 (2안)', '', '## ① 올리는 순서', '']
    d.add_heading('① 올리는 순서', level=2)
    for f in order:
        f = f.replace(os.sep, '/')
        p = d.add_paragraph(style='List Number'); r = p.add_run(f); r.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x4D, 0x3A)
        md.append(f'1. **[{f}]**')
    d.add_heading('② 맨 아래 검색용 텍스트', level=2)
    md += ['', '## ② 맨 아래 검색용 텍스트', '']
    for t in texts:
        if t['head']:
            d.add_heading(t['head'], level=3); md += ['### ' + t['head'], '']
        body = ' '.join(t['lines'])
        d.add_paragraph(body); md += [body, '']
    d.save(os.path.join(OUT, 'text.docx'))
    open(os.path.join(OUT, 'text.md'), 'w', encoding='utf-8').write('\n'.join(md))


if __name__ == '__main__':
    main()
